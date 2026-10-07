"""Concentration label colours, readability styles and overlap (meeting 1 Oct 2026)."""

from __future__ import annotations

import itertools
import zipfile
from io import BytesIO

import pandas as pd
import pytest

from models import Collar, EnvironmentalReading, Lithology, WaterLevel
from render_profiles import CONSULTING_SECTION_PROFILE, SECTION_SHEET_PROFILE
from render_theme import CHEMISTRY_FIXED_COLORS, parameter_series_color
from renderer import CrossSectionRenderer
from tests.conftest import run_pipeline


def test_label_colour_names_are_fixed_and_blue_is_reserved() -> None:
    base = dict(hole_id="BH-01", parameter="Chloride", value=5.0, depth=1.0)
    assert EnvironmentalReading(**base, label_color=" Green ").label_color == "green"
    assert EnvironmentalReading(**base, label_color=float("nan")).label_color == ""
    with pytest.raises(ValueError, match="reserved for groundwater"):
        EnvironmentalReading(**base, label_color="blue")
    with pytest.raises(ValueError, match="green, red, black or orange"):
        EnvironmentalReading(**base, label_color="yellow")
    assert set(CHEMISTRY_FIXED_COLORS) == {"green", "red", "black", "orange"}


def test_parameter_colour_is_stable_across_sections() -> None:
    """Colours used to follow the order parameters appeared on a section."""
    assert parameter_series_color("Chloride") == parameter_series_color(" chloride ")
    from render_theme import PARAMETER_PALETTE

    def is_blue(hex_colour: str) -> bool:
        r, g, b = (int(hex_colour[i : i + 2], 16) for i in (1, 3, 5))
        return b > g > r  # cobalt/sky blues (purple has r > g)

    assert not any(is_blue(c) for c in PARAMETER_PALETTE)  # blue is groundwater


def test_template_has_a_label_colour_drop_down_and_round_trips_it() -> None:
    from openpyxl import load_workbook

    from ingestion import ingest_workbook
    from workbook_template import (
        ENVIRONMENTAL_SHEET,
        build_input_template_bytes,
        export_cleaned_workbook_bytes,
    )

    book = load_workbook(BytesIO(build_input_template_bytes()))
    sheet = book[ENVIRONMENTAL_SHEET]
    header = [cell.value for cell in sheet[1]]
    assert "label_color" in header
    validations = sheet.data_validations.dataValidation
    assert any("green,red,black,orange" in str(v.formula1) for v in validations)

    collars = pd.DataFrame(
        [{"hole_id": "BH-01", "easting": 0, "northing": 0, "elevation": 100, "total_depth": 10}]
    )
    lith = pd.DataFrame([{"hole_id": "BH-01", "from_depth": 0, "to_depth": 10, "lithology_code": "Clay"}])
    env = pd.DataFrame(
        [
            {"hole_id": "BH-01", "parameter": "Chloride", "value": 12.0, "depth": 2.0, "label_color": "red"},
            {"hole_id": "BH-01", "parameter": "Chloride", "value": 1.0, "depth": 4.0, "label_color": "blue"},
        ]
    )
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        collars.to_excel(writer, sheet_name="Collars", index=False)
        lith.to_excel(writer, sheet_name="Lithology", index=False)
        env.to_excel(writer, sheet_name="Environmental", index=False)
    buffer.seek(0)
    result, _ = ingest_workbook(buffer)
    assert [r.label_color for r in result.environmental_readings] == ["red"]
    assert any("Environmental row 3" in e and "reserved" in e for e in result.errors)

    exported = export_cleaned_workbook_bytes(result, project_metadata={})
    again, _ = ingest_workbook(BytesIO(exported))
    assert again.environmental_readings[0].label_color == "red"
    assert zipfile.is_zipfile(BytesIO(exported))


def _chemistry_section(n_readings: int = 8):
    ids = ["BH-01", "BH-02", "BH-03"]
    collars = [Collar(hole_id=h, easting=i * 6.0, northing=0.0, elevation=100.0, total_depth=12.0) for i, h in enumerate(ids)]
    lith = [Lithology(hole_id=h, from_depth=0.0, to_depth=12.0, lithology_code="Sandy Clay") for h in ids]
    readings = [
        EnvironmentalReading(
            hole_id=h, parameter="Chloride", value=float(100 + k * 37), depth=0.5 + k * (11.0 / n_readings),
            label_color=["green", "red", "orange", "black", ""][k % 5],
        )
        for h in ids
        for k in range(n_readings)
    ]
    water = [WaterLevel(hole_id=h, depth=3.0 + i) for i, h in enumerate(ids)]
    projected, polygons, _ = run_pipeline(collars, lith, [(0.0, 0.0), (12.0, 0.0)])
    return ids, projected, polygons, readings, water


def _render(style: str, profile=SECTION_SHEET_PROFILE, **extra):
    ids, projected, polygons, readings, water = _chemistry_section()
    renderer = CrossSectionRenderer(
        show_legend=True,
        render_profile=profile.model_copy(
            update={
                "show_parameter_markers": True,
                "show_parameter_labels": True,
                "chemistry_label_style": style,
                "show_water_elevation_labels": True,
                **extra,
            }
        ),
        environmental_readings=readings,
        environmental_parameters=("Chloride",),
    )
    figure = renderer.render(polygons, projected, collar_depths={h: 12.0 for h in ids}, water_levels=water)
    return renderer, figure


def test_workbook_colours_reach_the_figure_and_labels_never_overlap() -> None:
    from matplotlib.text import Text

    from renderer_water import _figure_renderer, _overlap_area

    renderer, figure = _render("plain")
    mpl_renderer = _figure_renderer(figure)
    figure.draw_without_rendering()
    labels = [(kind, ann) for kind, ann, _c in renderer._water_labels]
    kinds = {kind for kind, _ in labels}
    assert "chem" in kinds and "rl" in kinds
    colours = {ann.get_color().lower() for kind, ann in labels if kind == "chem"}
    assert {CHEMISTRY_FIXED_COLORS["green"].lower(), CHEMISTRY_FIXED_COLORS["red"].lower()} <= colours
    boxes = [Text.get_window_extent(ann, mpl_renderer) for _, ann in labels]
    overlaps = sum(1 for a, b in itertools.combinations(boxes, 2) if _overlap_area(a, b) > 0)
    assert overlaps == 0
    # every label stays on the page
    assert all(box.x0 >= figure.bbox.x0 and box.x1 <= figure.bbox.x1 for box in boxes)


def test_readability_styles_change_the_label_rendering() -> None:
    _, plain = _render("plain")
    _, boxed = _render("box")
    _, dotted = _render("dot")
    _, stroked = _render("stroke")

    def chem_texts(fig):
        return [t for ax in fig.axes for t in ax.texts if t.get_text().strip().isdigit()]

    # Plain keeps the translucent marker-mode box; "box" is fully opaque.
    assert all((t.get_bbox_patch() is None or t.get_bbox_patch().get_alpha() < 1.0) for t in chem_texts(plain))
    assert all(t.get_bbox_patch() is not None and t.get_bbox_patch().get_alpha() == 1.0 for t in chem_texts(boxed))
    # The label stays real (searchable) text; a white twin beneath it carries
    # the outline path effect.
    stroked_labels = [t for t in chem_texts(stroked) if t.get_color() != "white"]
    halos = [t for t in chem_texts(stroked) if t.get_color() == "white"]
    assert stroked_labels and not any(t.get_path_effects() for t in stroked_labels)
    assert len(halos) == len(stroked_labels) and all(t.get_path_effects() for t in halos)
    assert all(t.get_color().lower() == "#111827" for t in chem_texts(dotted))  # black text
    # With series markers drawn, the colour rides on the column marker (no
    # second dot); the scatter carries per-reading colours.
    from matplotlib.colors import to_hex

    scatter_colours = {
        to_hex(c).upper()
        for ax in dotted.axes
        for coll in ax.collections
        if hasattr(coll, "get_offsets") and len(coll.get_offsets()) > 1
        for c in coll.get_facecolors()
    }
    assert {CHEMISTRY_FIXED_COLORS["green"], CHEMISTRY_FIXED_COLORS["red"]} <= scatter_colours



def test_consulting_layout_also_registers_chemistry_labels() -> None:
    renderer, _ = _render("box", profile=CONSULTING_SECTION_PROFILE)
    assert any(kind == "chem" for kind, _a, _c in renderer._water_labels)


def test_colour_header_accepts_uk_spelling_and_drop_down_blocks_typed_values() -> None:
    from openpyxl import load_workbook

    from ingestion import ingest_workbook
    from workbook_template import (
        ENVIRONMENTAL_SHEET,
        build_input_template_bytes,
        export_cleaned_workbook_bytes,
    )

    collars = pd.DataFrame([{"hole_id": "BH-01", "easting": 0, "northing": 0, "elevation": 100, "total_depth": 10}])
    lith = pd.DataFrame([{"hole_id": "BH-01", "from_depth": 0, "to_depth": 10, "lithology_code": "Clay"}])
    env = pd.DataFrame([{"hole_id": "BH-01", "parameter": "Chloride", "value": 12.0, "depth": 2.0, "Label Colour": "orange"}])
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        collars.to_excel(writer, sheet_name="Collars", index=False)
        lith.to_excel(writer, sheet_name="Lithology", index=False)
        env.to_excel(writer, sheet_name="Environmental", index=False)
    buffer.seek(0)
    result, _ = ingest_workbook(buffer)
    assert result.environmental_readings[0].label_color == "orange"

    def validations(book_bytes):
        sheet = load_workbook(BytesIO(book_bytes))[ENVIRONMENTAL_SHEET]
        return [v for v in sheet.data_validations.dataValidation if "green,red" in str(v.formula1)]

    template = validations(build_input_template_bytes())
    assert template and all(v.showErrorMessage for v in template)  # typed "yellow" is rejected in Excel
    exported = validations(export_cleaned_workbook_bytes(result, project_metadata={}))
    assert exported and all(v.showErrorMessage for v in exported)


def test_parameters_on_one_section_never_share_a_colour_and_threshold_middle_is_orange() -> None:
    from render_theme import (
        CHEMISTRY_LABEL_ORANGE,
        PARAMETER_PALETTE,
        chemistry_label_color,
        parameter_series_colors,
    )

    names = ["Chloride", "BTEX", "Sulphate", "Benzene", "pH"]
    assigned = parameter_series_colors(names)
    assert len(set(assigned.values())) == len(names)
    assert set(assigned.values()) <= set(PARAMETER_PALETTE)
    # Parameters whose hashed slots do not collide keep their stable colour.
    stable = parameter_series_colors(["Benzene", "Xylenes", "Toluene"])
    assert all(stable[name] == parameter_series_color(name) for name in stable)
    assert chemistry_label_color(150.0, "threshold", green_max=100.0, yellow_max=250.0) == CHEMISTRY_LABEL_ORANGE
    assert "#CA8A04" not in {  # the old yellow is gone from threshold mode
        chemistry_label_color(v, "threshold", green_max=100.0, yellow_max=250.0) for v in (1.0, 150.0, 999.0)
    }


def _dense_two_hole_section(n_readings: int = 15):
    """Two holes 0.5 m apart, dense readings: the case where labels used to
    migrate onto the neighbouring column with no leader."""
    ids = ["BH-01", "BH-02", "BH-03"]
    collars = [
        Collar(hole_id="BH-01", easting=0.0, northing=0.0, elevation=100.0, total_depth=12.0),
        Collar(hole_id="BH-02", easting=0.5, northing=0.0, elevation=100.0, total_depth=12.0),
        Collar(hole_id="BH-03", easting=20.0, northing=0.0, elevation=100.0, total_depth=12.0),
    ]
    lith = [Lithology(hole_id=h, from_depth=0.0, to_depth=12.0, lithology_code="Clay") for h in ids]
    readings = [
        EnvironmentalReading(hole_id=h, parameter="Chloride", value=float(100 + k * 37), depth=0.3 + k * 0.75)
        for h in ids
        for k in range(n_readings)
    ]
    projected, polygons, _ = run_pipeline(collars, lith, [(0.0, 0.0), (20.0, 0.0)])
    return ids, projected, polygons, readings


@pytest.mark.parametrize("profile", [SECTION_SHEET_PROFILE, CONSULTING_SECTION_PROFILE])
def test_value_labels_never_sit_over_any_borehole_column(profile) -> None:
    from matplotlib.text import Text

    from renderer_water import _figure_renderer, _overlap_area

    ids, projected, polygons, readings = _dense_two_hole_section()
    renderer = CrossSectionRenderer(
        show_legend=True,
        render_profile=profile.model_copy(
            update={"show_parameter_markers": True, "show_parameter_labels": True, "parameter_draw_leaders": False}
        ),
        environmental_readings=readings,
        environmental_parameters=("Chloride",),
    )
    figure = renderer.render(polygons, projected, collar_depths={h: 12.0 for h in ids})
    mpl_renderer = _figure_renderer(figure)
    figure.draw_without_rendering()
    columns = renderer._column_obstacle_boxes(figure)
    assert len(columns) == 3
    labels = [ann for kind, ann, _c in renderer._water_labels if kind == "chem"]
    assert len(labels) == 45
    boxes = [Text.get_window_extent(ann, mpl_renderer) for ann in labels]
    over_columns = sum(1 for box in boxes for col in columns if _overlap_area(box, col) > 0)
    assert over_columns == 0
    assert all(ann.get_horizontalalignment() == "left" for ann in labels)  # never flipped leftwards
    overlaps = sum(1 for a, b in itertools.combinations(boxes, 2) if _overlap_area(a, b) > 0)
    assert overlaps == 0


def test_dot_style_dot_follows_its_label_when_moved() -> None:
    from matplotlib.text import Text

    from renderer_water import _figure_renderer

    ids, projected, polygons, readings = _dense_two_hole_section()
    renderer = CrossSectionRenderer(
        show_legend=True,
        render_profile=SECTION_SHEET_PROFILE.model_copy(
            update={
                "show_parameter_markers": True,
                "show_parameter_labels": True,
                "chemistry_label_style": "dot",
                "parameter_draw_markers": False,  # label dots are used only without series markers
            }
        ),
        environmental_readings=readings,
        environmental_parameters=("Chloride",),
    )
    figure = renderer.render(polygons, projected, collar_depths={h: 12.0 for h in ids})
    mpl_renderer = _figure_renderer(figure)
    figure.draw_without_rendering()
    moved = 0
    for kind, ann, _c in renderer._water_labels:
        if kind != "chem":
            continue
        dot = ann._chem_dot
        text_box = Text.get_window_extent(ann, mpl_renderer)
        dot_xy = dot.get_transform().transform([[dot.get_xdata()[0], dot.get_ydata()[0]]])[0]
        # Dot sits just left of the text, vertically centred on it.
        assert text_box.x0 - 12 <= dot_xy[0] <= text_box.x0 + 1
        assert abs(dot_xy[1] - (text_box.y0 + text_box.y1) / 2) < 3
        if abs(ann.xyann[1]) > 14:
            moved += 1
    assert moved > 0  # the scenario really exercised the collision pass


def test_outlined_labels_remain_searchable_text_in_pdf() -> None:
    from pypdf import PdfReader

    from export_framing import ExportFramingConfig

    ids, projected, polygons, readings, water = _chemistry_section(n_readings=3)
    renderer = CrossSectionRenderer(
        show_legend=True,
        render_profile=SECTION_SHEET_PROFILE.model_copy(
            update={"show_parameter_markers": True, "show_parameter_labels": True, "chemistry_label_style": "stroke"}
        ),
        environmental_readings=readings,
        environmental_parameters=("Chloride",),
        export_framing=ExportFramingConfig(page_preset="letter_landscape", export_dpi=72),
    )
    figure = renderer.render(polygons, projected, collar_depths={h: 12.0 for h in ids}, water_levels=water)
    text = PdfReader(BytesIO(renderer.to_pdf_bytes(figure))).pages[0].extract_text()
    values = {f"{r.value:g}" for r in readings}
    assert sum(1 for v in values if v in text) >= len(values) - 1


@pytest.mark.parametrize("style", ["plain", "stroke", "box", "dot"])
def test_threshold_key_is_added_without_touching_the_collision_pass(style) -> None:
    from matplotlib.offsetbox import AnnotationBbox

    black, _ = _render(style)
    threshold, figure = _render(
        style,
        chemistry_color_mode="threshold",
        chemistry_threshold_green_max=5.0,
        chemistry_threshold_yellow_max=10.0,
    )
    assert getattr(black, "chemistry_threshold_key_text", None) is None
    assert threshold.chemistry_threshold_key_text.startswith("green ≤ 5 · orange 5–10")
    keys = [a for ax in figure.axes for a in ax.artists if isinstance(a, AnnotationBbox)]
    assert len(keys) == 1 and keys[0].get_gid() == "chemistry-threshold-key"
    # The key is a fixed panel, not a value label the collision pass moves.
    assert len(threshold._water_labels) == len(black._water_labels)
    import matplotlib.pyplot as plt

    plt.close("all")


def _p2_render(mode: str, *, label_colors=("", "green", "")):
    """Consulting sheet like the P2 "Chemistry columns" style (no markers)."""
    ids = ["BH-01", "BH-02"]
    collars = [
        Collar(hole_id=h, easting=i * 8.0, northing=0.0, elevation=100.0, total_depth=10.0)
        for i, h in enumerate(ids)
    ]
    lith = [Lithology(hole_id=h, from_depth=0.0, to_depth=10.0, lithology_code="Sandy Clay") for h in ids]
    readings = [
        EnvironmentalReading(
            hole_id=h, parameter="Chloride", value=float(120 + 50 * k + 7 * i), depth=1.0 + 3.0 * k,
            unit="mg/kg", label_color=colour,
        )
        for i, h in enumerate(ids)
        for k, colour in enumerate(label_colors)
    ]
    projected, polygons, _ = run_pipeline(collars, lith, [(0.0, 0.0), (8.0, 0.0)])
    renderer = CrossSectionRenderer(
        render_profile=CONSULTING_SECTION_PROFILE.model_copy(
            update={
                "show_parameter_markers": True,
                "show_parameter_labels": True,
                "parameter_draw_markers": False,
                "show_track_lithology": True,
                "chemistry_color_mode": mode,
            }
        ),
        environmental_readings=readings,
        environmental_parameters=("Chloride",),
    )
    renderer.render(polygons, projected, collar_depths={h: 10.0 for h in ids})
    colours = {
        ann.get_text(): ann.get_color().upper()
        for kind, ann, _c in renderer._water_labels
        if kind == "chem"
    }
    return renderer, colours


def test_p2_chemistry_columns_preset_prints_values_red_with_legend_sample() -> None:
    import matplotlib.pyplot as plt

    from render_theme import CHEMISTRY_LABEL_BLACK, CHEMISTRY_LABEL_RED, chemistry_label_color
    from ui_output_presets import OUTPUT_PRESETS

    mode = OUTPUT_PRESETS["p2_chemistry_sticks"].chemistry_color_mode
    assert mode == "red"
    # Client figures print pure #FF0000 (4.0:1); the agreed red clears AA.
    assert CHEMISTRY_LABEL_RED == CHEMISTRY_FIXED_COLORS["red"]
    assert chemistry_label_color(999.0, "red") == CHEMISTRY_LABEL_RED
    renderer, colours = _p2_render(mode)
    assert colours["120"] == CHEMISTRY_LABEL_RED and colours["227"] == CHEMISTRY_LABEL_RED
    # A colour picked in the workbook still wins over the preset default.
    assert colours["170"] == CHEMISTRY_FIXED_COLORS["green"].upper()
    assert colours["177"] == CHEMISTRY_FIXED_COLORS["green"].upper()
    (entry,) = renderer.parameter_series_legend
    assert entry["label"] == "CHLORIDE CONCENTRATION (mg/kg)"
    assert entry["color"] == CHEMISTRY_LABEL_BLACK  # label text stays black
    assert entry["sample_text"] == "120" and entry["sample_color"] == CHEMISTRY_LABEL_RED
    # Legend panels draw the sample through the chemistry mixin helper.
    fig, ax = plt.subplots()
    artist = renderer._draw_parameter_legend_sample(ax, 0.1, 0.5, entry, font_size=7.0)
    assert artist.get_text() == "120" and artist.get_color() == CHEMISTRY_LABEL_RED
    blank = {**entry, "sample_text": ""}
    assert renderer._draw_parameter_legend_sample(ax, 0.1, 0.5, blank, font_size=7.0) is None
    plt.close("all")


def test_black_default_and_threshold_modes_unchanged_without_legend_sample() -> None:
    import matplotlib.pyplot as plt

    from render_theme import (
        CHEMISTRY_LABEL_BLACK,
        CHEMISTRY_LABEL_GREEN,
        CHEMISTRY_LABEL_RED,
        chemistry_label_color,
    )
    from ui_output_presets import OUTPUT_PRESETS

    for preset_id, preset in OUTPUT_PRESETS.items():
        if preset_id != "p2_chemistry_sticks":
            assert preset.chemistry_color_mode in (None, "black")
    renderer, colours = _p2_render("black")
    assert colours["120"] == CHEMISTRY_LABEL_BLACK
    assert colours["170"] == CHEMISTRY_FIXED_COLORS["green"].upper()
    assert "sample_text" not in renderer.parameter_series_legend[0]

    threshold, colours = _p2_render("threshold")  # no limits set: black, no key
    assert colours["120"] == CHEMISTRY_LABEL_BLACK
    assert threshold.chemistry_threshold_key_text is None
    assert "sample_text" not in threshold.parameter_series_legend[0]
    bands = {"green_max": 150.0, "yellow_max": 200.0}
    assert chemistry_label_color(120.0, "threshold", **bands) == CHEMISTRY_LABEL_GREEN
    assert chemistry_label_color(270.0, "threshold", **bands) == CHEMISTRY_LABEL_RED
    plt.close("all")


def _threshold_key_figure(profile, *, page=None, font_size=None, mode="threshold"):
    from export_framing import ExportFramingConfig
    from models import SectionFigureMetadata

    ids, projected, polygons, readings, water = _chemistry_section()
    update = {
        "show_parameter_markers": True,
        "show_parameter_labels": True,
        "chemistry_color_mode": mode,
        "chemistry_threshold_green_max": 150.0,
        "chemistry_threshold_yellow_max": 250.0,
    }
    if font_size is not None:
        update["export_font_size"] = font_size
    renderer = CrossSectionRenderer(
        show_legend=True,
        render_profile=profile.model_copy(update=update),
        environmental_readings=readings,
        environmental_parameters=("Chloride",),
        disclaimer="Interpreted between boreholes.",
        figure_metadata=SectionFigureMetadata(coordinate_reference="EPSG:26912", elevation_datum="CGVD2013"),
        export_framing=ExportFramingConfig(page_preset=page, export_dpi=72) if page else None,
    )
    figure = renderer.render(polygons, projected, collar_depths={h: 12.0 for h in ids}, water_levels=water)
    if page:
        renderer.to_png_bytes(figure, dpi=72)
    return renderer, figure


def _threshold_key_artists(figure):
    from matplotlib.offsetbox import AnnotationBbox

    return [
        a for ax in figure.axes for a in ax.artists
        if isinstance(a, AnnotationBbox) and a.get_gid() == "chemistry-threshold-key"
    ]


def _assert_key_off_plot_and_clear(figure) -> None:
    from matplotlib.text import Text

    from renderer_water import _figure_renderer, _overlap_area

    mpl_renderer = _figure_renderer(figure)
    figure.draw_without_rendering()
    (key,) = _threshold_key_artists(figure)
    key_box = key.get_window_extent(mpl_renderer)
    main = key.axes
    assert key_box.y1 <= main.get_window_extent(mpl_renderer).y0  # off the plot
    assert key_box.y1 <= main.xaxis.get_tightbbox(mpl_renderer).y0 + 1e-6  # under ticks + label
    supx = getattr(figure, "_supxlabel", None)
    if supx is not None and supx.get_text():
        assert key_box.y1 <= supx.get_window_extent(mpl_renderer).y0 + 1e-6
    page = figure.bbox
    assert page.x0 <= key_box.x0 and key_box.x1 <= page.x1 and page.y0 <= key_box.y0
    own = {id(t) for t in key.offsetbox.findobj(Text)}
    for ax in figure.axes:
        if not ax.axison:  # hidden axis furniture of panel axes is never drawn
            own |= {id(t) for t in ax.xaxis.findobj(Text) + ax.yaxis.findobj(Text)}
    for text in figure.findobj(Text):
        if id(text) in own or not text.get_visible() or not text.get_text().strip():
            continue
        assert _overlap_area(key_box, text.get_window_extent(mpl_renderer)) <= 0, text.get_text()
    main_box = main.get_window_extent(mpl_renderer)
    for ax in figure.axes:
        box = ax.get_window_extent(mpl_renderer)
        if box.y1 < main_box.y0:  # consulting subtitle band / title block panels
            assert _overlap_area(key_box, box) <= 0


@pytest.mark.parametrize(
    "profile_name,show_headers",
    [("consulting", None), ("sheet", True), ("sheet", False), ("chart", None)],
)
@pytest.mark.parametrize(
    "page,font_size",
    [
        (None, None),
        ("letter_landscape", 8),
        ("letter_portrait", None),
        ("letter_portrait", 14),
        ("tabloid_landscape", 14),
    ],
)
def test_threshold_key_sits_below_the_distance_label_clear_of_everything(
    profile_name, show_headers, page, font_size
) -> None:
    import matplotlib.pyplot as plt

    from render_profiles import CHART_PROFILE

    profile = {
        "consulting": CONSULTING_SECTION_PROFILE,
        "sheet": SECTION_SHEET_PROFILE,
        "chart": CHART_PROFILE,
    }[profile_name]
    if show_headers is not None:
        profile = profile.model_copy(update={"show_column_headers": show_headers})
    _, figure = _threshold_key_figure(profile, page=page, font_size=font_size)
    _assert_key_off_plot_and_clear(figure)
    plt.close("all")


@pytest.mark.parametrize("mode", ["black", "red"])
def test_threshold_key_absent_in_fixed_colour_modes(mode) -> None:
    import matplotlib.pyplot as plt

    for profile in (CONSULTING_SECTION_PROFILE, SECTION_SHEET_PROFILE):
        renderer, figure = _threshold_key_figure(profile, mode=mode)
        assert renderer.chemistry_threshold_key_text is None
        assert _threshold_key_artists(figure) == []
    plt.close("all")


def test_consulting_band_only_moves_for_the_threshold_key() -> None:
    import matplotlib.pyplot as plt

    _, black = _threshold_key_figure(CONSULTING_SECTION_PROFILE, mode="black")
    _, keyed = _threshold_key_figure(CONSULTING_SECTION_PROFILE)
    black_band = [ax.get_position().bounds for ax in black.axes[1:5]]
    keyed_band = [ax.get_position().bounds for ax in keyed.axes[1:5]]
    # Main plot and title block stay put; only the subtitle band yields room.
    assert black.axes[0].get_position().bounds == keyed.axes[0].get_position().bounds
    assert black_band[3] == keyed_band[3]
    for before, after in zip(black_band[:3], keyed_band[:3], strict=True):
        assert after[1] + after[3] < before[1] + before[3]  # band top lowered
        assert after[3] > 0.6 * before[3]  # only slightly
    plt.close("all")
