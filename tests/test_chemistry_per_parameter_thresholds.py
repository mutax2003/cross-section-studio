"""Per-parameter green / orange / red threshold limits (client meeting 8 Oct 2026).

Each parameter plotted on a section is banded by its own limits in its own
units; the sheet-wide pair is only the fallback for parameters without one.
"""

from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
import pytest

from models import Collar, EnvironmentalReading, Lithology
from render_profiles import (
    CONSULTING_SECTION_PROFILE,
    SECTION_SHEET_PROFILE,
    CrossSectionRenderProfile,
    chemistry_thresholds_from_readings,
    normalise_chemistry_thresholds,
    resolved_chemistry_thresholds,
    resolved_chemistry_thresholds_for,
)
from render_theme import CHEMISTRY_LABEL_GREEN, CHEMISTRY_LABEL_ORANGE, CHEMISTRY_LABEL_RED
from renderer import CrossSectionRenderer
from section_build_request import SectionBuildRequest
from tests.conftest import run_pipeline

ROOT = Path(__file__).resolve().parents[1]

THRESHOLDS = {"Chloride": (100.0, 250.0), "Benzene": (0.5, 5.0)}


# --- resolver ------------------------------------------------------------------


def test_normalise_accepts_pairs_lists_and_mappings_and_drops_blank_rows() -> None:
    cleaned = normalise_chemistry_thresholds(
        {
            " Chloride ": (100, 250),
            "benzene": {"green_max": "0.5", "yellow_max": None},
            "TPH": [None, None],
            "": (1, 2),
            "Arsenic": {"green_max": 2, "orange_max": 10},
        }
    )
    assert cleaned == {"Chloride": (100.0, 250.0), "benzene": (0.5, None), "Arsenic": (2.0, 10.0)}
    with pytest.raises(ValueError, match=">= 0"):
        normalise_chemistry_thresholds({"Chloride": (-1, 5)})


def test_per_parameter_resolver_falls_back_limit_by_limit() -> None:
    # Own entry wins, case-insensitively.
    assert resolved_chemistry_thresholds_for("chloride", THRESHOLDS, 7.0, 9.0) == (100.0, 250.0)
    # A parameter without an entry takes the sheet-wide pair.
    assert resolved_chemistry_thresholds_for("TPH", THRESHOLDS, 7.0, 9.0) == (7.0, 9.0)
    # ...and the Configure defaults when that pair is missing too.
    assert resolved_chemistry_thresholds_for("TPH", THRESHOLDS) == resolved_chemistry_thresholds(None, None)
    # One blank limit in the parameter's own entry is filled from the pair.
    partial = {"Benzene": (0.5, None)}
    assert resolved_chemistry_thresholds_for("Benzene", partial, 7.0, 9.0) == (0.5, 9.0)
    # The no-inversion rule still applies after the fill-in (orange never below green).
    assert resolved_chemistry_thresholds_for("Benzene", partial, 7.0, 0.2) == (0.2, 0.2)
    # Empty map behaves exactly like the old global resolver.
    assert resolved_chemistry_thresholds_for("x", {}, 3.0, None) == resolved_chemistry_thresholds(3.0, None)


def test_profile_and_request_keep_the_global_pair_as_fallback() -> None:
    profile = CrossSectionRenderProfile(chemistry_thresholds=THRESHOLDS)
    assert profile.chemistry_thresholds == THRESHOLDS
    assert profile.chemistry_threshold_green_max is None  # the old pair is untouched
    # Existing kwargs keep working without the map.
    assert CrossSectionRenderProfile(chemistry_threshold_green_max=5.0).chemistry_thresholds == {}


# --- request / cache ----------------------------------------------------------------


def _request(**extra) -> SectionBuildRequest:
    return SectionBuildRequest(transect_points=((0.0, 0.0), (10.0, 0.0)), **extra)


def test_request_json_round_trip_and_geometry_cache_exclusion() -> None:
    request = _request(chemistry_color_mode="threshold", chemistry_thresholds=THRESHOLDS)
    dumped = request.model_dump(mode="json")
    assert dumped["chemistry_thresholds"] == {"Chloride": [100.0, 250.0], "Benzene": [0.5, 5.0]}
    again = SectionBuildRequest.model_validate(json.loads(json.dumps(dumped)))
    assert again == request
    assert again.chemistry_thresholds["Benzene"] == (0.5, 5.0)
    # Cosmetic: never part of the geometry cache key, always part of the full key.
    plain = _request()
    assert "chemistry_thresholds" not in json.dumps(request.geometry_cache_payload())
    assert request.geometry_cache_key(("BH-01",)) == plain.geometry_cache_key(("BH-01",))
    assert request.cache_key(("BH-01",)) != plain.cache_key(("BH-01",))
    # Blank entries are normalised away on construction.
    assert _request(chemistry_thresholds={"TPH": (None, None)}).chemistry_thresholds == {}


# --- renderer -----------------------------------------------------------------------


def _two_parameter_section():
    ids = ["BH-01", "BH-02", "BH-03"]
    collars = [
        Collar(hole_id=h, easting=i * 30.0, northing=0.0, elevation=100.0, total_depth=12.0)
        for i, h in enumerate(ids)
    ]
    lith = [Lithology(hole_id=h, from_depth=0.0, to_depth=12.0, lithology_code="Sandy Clay") for h in ids]
    # Same depths: chloride 40 / 180 / 620 mg/kg, benzene 0.2 / 1.4 / 7.5 mg/kg.
    chloride = (40.0, 180.0, 620.0)
    benzene = (0.2, 1.4, 7.5)
    readings = []
    for i, h in enumerate(ids):
        readings.append(EnvironmentalReading(hole_id=h, parameter="Chloride", value=chloride[i], depth=2.0, unit="mg/kg"))
        readings.append(EnvironmentalReading(hole_id=h, parameter="Benzene", value=benzene[i], depth=7.0, unit="mg/kg"))
    projected, polygons, _ = run_pipeline(collars, lith, [(0.0, 0.0), (60.0, 0.0)])
    return ids, projected, polygons, readings


def _render(profile, parameters=("Chloride", "Benzene"), **profile_updates):
    ids, projected, polygons, readings = _two_parameter_section()
    renderer = CrossSectionRenderer(
        show_legend=True,
        render_profile=profile.model_copy(
            update={
                "show_parameter_markers": True,
                "show_parameter_labels": True,
                "chemistry_color_mode": "threshold",
                **profile_updates,
            }
        ),
        environmental_readings=readings,
        environmental_parameters=parameters,
    )
    figure = renderer.render(polygons, projected, collar_depths={h: 12.0 for h in ids})
    return renderer, figure


def _chem_label_colours(renderer) -> dict[str, str]:
    return {
        ann.get_text(): ann.get_color().lower()
        for kind, ann, _c in renderer._water_labels
        if kind == "chem"
    }


@pytest.mark.parametrize("profile", [SECTION_SHEET_PROFILE, CONSULTING_SECTION_PROFILE])
def test_two_parameters_are_coloured_by_their_own_limits(profile) -> None:
    renderer, figure = _render(profile, chemistry_thresholds=THRESHOLDS)
    try:
        colours = _chem_label_colours(renderer)
        green, orange, red = (c.lower() for c in (CHEMISTRY_LABEL_GREEN, CHEMISTRY_LABEL_ORANGE, CHEMISTRY_LABEL_RED))
        assert colours["40"] == green and colours["180"] == orange and colours["620"] == red
        # Under the chloride limits every benzene value would be green; its own
        # 0.5 / 5 limits band them green / orange / red.
        assert colours["0.2"] == green and colours["1.4"] == orange and colours["7.5"] == red
        # Dots agree with their labels (one scatter per parameter, in sorted order).
        dots = [c for ax in figure.axes for c in ax.collections if getattr(c, "_css_marker_kind", None) == "chem"]
        assert len(dots) >= 2
        # Key rows in Configure order (the order passed in), not alphabetical.
        assert renderer.chemistry_threshold_key_lines == (
            "Chloride: green ≤ 100 · orange 100–250 · red > 250 mg/kg",
            "Benzene: green ≤ 0.5 · orange 0.5–5 · red > 5 mg/kg",
        )
        assert renderer.chemistry_threshold_key_text == "\n".join(renderer.chemistry_threshold_key_lines)
        # Legend 3-dot sample per parameter: benzene has an orange band, so three dots.
        by_name = {entry["parameter"]: entry for entry in renderer.parameter_series_legend}
        assert len(by_name["Benzene"]["threshold_colors"]) == 3
    finally:
        plt.close(figure)


def test_global_pair_colours_a_parameter_without_its_own_entry() -> None:
    renderer, figure = _render(
        SECTION_SHEET_PROFILE,
        chemistry_thresholds={"Benzene": (0.5, 5.0)},
        chemistry_threshold_green_max=50.0,
        chemistry_threshold_yellow_max=200.0,
    )
    try:
        colours = _chem_label_colours(renderer)
        # Chloride falls back to the sheet-wide 50 / 200 pair.
        assert colours["40"] == CHEMISTRY_LABEL_GREEN.lower()
        assert colours["180"] == CHEMISTRY_LABEL_ORANGE.lower()
        assert colours["620"] == CHEMISTRY_LABEL_RED.lower()
        assert colours["1.4"] == CHEMISTRY_LABEL_ORANGE.lower()
        assert "Chloride: green ≤ 50 · orange 50–200 · red > 200 mg/kg" in (
            renderer.chemistry_threshold_key_lines
        )
    finally:
        plt.close(figure)


def test_single_parameter_key_keeps_the_bare_band_text() -> None:
    renderer, figure = _render(SECTION_SHEET_PROFILE, parameters=("Benzene",), chemistry_thresholds=THRESHOLDS)
    try:
        assert renderer.chemistry_threshold_key_text == "green ≤ 0.5 · orange 0.5–5 · red > 5 mg/kg"
        assert len(renderer.chemistry_threshold_key_lines) == 1
    finally:
        plt.close(figure)


def test_empty_orange_band_is_omitted_per_parameter() -> None:
    renderer, figure = _render(
        SECTION_SHEET_PROFILE, chemistry_thresholds={"Chloride": (100.0, 250.0), "Benzene": (1.0, 1.0)}
    )
    try:
        assert renderer.chemistry_threshold_key_lines[1] == "Benzene: green ≤ 1 · red > 1 mg/kg"
        by_name = {entry["parameter"]: entry for entry in renderer.parameter_series_legend}
        assert len(by_name["Benzene"]["threshold_colors"]) == 2
        assert len(by_name["Chloride"]["threshold_colors"]) == 3
    finally:
        plt.close(figure)


def test_pipeline_threads_the_map_to_the_figure() -> None:
    from pipeline import build_cross_section

    ids, _projected, _polygons, readings = _two_parameter_section()
    collars = [Collar(hole_id=h, easting=i * 30.0, northing=0.0, elevation=100.0, total_depth=12.0) for i, h in enumerate(ids)]
    lith = [Lithology(hole_id=h, from_depth=0.0, to_depth=12.0, lithology_code="Sandy Clay") for h in ids]
    result = build_cross_section(
        collars,
        lith,
        [(0.0, 0.0), (60.0, 0.0)],
        environmental_readings=readings,
        environmental_parameters=("Chloride", "Benzene"),
        chemistry_color_mode="threshold",
        chemistry_thresholds=THRESHOLDS,
        export_formats=frozenset({"svg"}),
    )
    svg = result.svg_bytes.decode("utf-8")
    assert "Benzene:" in svg and "Chloride:" in svg


# --- workbook seeding -----------------------------------------------------------------


def _two_parameter_workbook_bytes(section_label: str = "A-A'") -> bytes:
    import openpyxl

    from workbook_template import (
        COLLAR_COLUMNS,
        ENVIRONMENTAL_COLUMNS,
        LITHOLOGY_COLUMNS,
        SECTION_COLUMNS,
        build_input_template_bytes,
    )

    # Limits stay just before the optional sampling ``date`` (appended last).
    assert ("green_max", "orange_max", "date") == ENVIRONMENTAL_COLUMNS[-3:]
    book = openpyxl.load_workbook(BytesIO(build_input_template_bytes()))

    def replace(sheet, columns, rows):
        header = [cell.value for cell in sheet[1][: len(columns)]]
        assert tuple(header) == columns, (sheet.title, header)
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
        for row_index, row in enumerate(rows, start=2):
            for col_index, column in enumerate(columns, start=1):
                value = row.get(column, "")
                sheet.cell(row=row_index, column=col_index, value=None if value == "" else value)

    holes = ["BH-01", "BH-02", "BH-03"]
    replace(
        book["Collars"],
        COLLAR_COLUMNS,
        [
            {"hole_id": h, "easting": 500000.0 + i * 30.0, "northing": 6000000.0, "elevation": 100.0, "total_depth": 12.0}
            for i, h in enumerate(holes)
        ],
    )
    replace(
        book["Lithology"],
        LITHOLOGY_COLUMNS,
        [
            {"hole_id": h, "from_depth": 0.0, "to_depth": 6.0, "lithology_code": "Clay", "unit_order": 1}
            for h in holes
        ]
        + [
            {"hole_id": h, "from_depth": 6.0, "to_depth": 12.0, "lithology_code": "Sand", "unit_order": 2}
            for h in holes
        ],
    )
    env = []
    for i, h in enumerate(holes):
        env.append(
            {
                "hole_id": h, "parameter": "Chloride", "value": (40.0, 180.0, 620.0)[i], "depth": 2.0,
                "unit": "mg/kg",
            }
        )
        env.append(
            {
                "hole_id": h, "parameter": "Benzene", "value": (0.2, 1.4, 7.5)[i], "depth": 7.0,
                "unit": "mg/kg",
                # One filled row per parameter is enough to seed the editor.
                "green_max": 0.5 if i == 0 else "",
                "orange_max": 5.0 if i == 0 else "",
            }
        )
    replace(book["Environmental"], ENVIRONMENTAL_COLUMNS, env)
    for name in ("Water", "Screens", "Gradients"):
        sheet = book[name]
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)
    replace(book["Sections"], SECTION_COLUMNS, [{"section_label": section_label, "hole_ids": ", ".join(holes)}])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_workbook_green_max_orange_max_seed_per_parameter_limits() -> None:
    from ingestion import ingest_workbook
    from workbook_template import export_cleaned_workbook_bytes

    result, _ = ingest_workbook(BytesIO(_two_parameter_workbook_bytes()))
    assert not result.errors
    assert chemistry_thresholds_from_readings(result.environmental_readings) == {"Benzene": (0.5, 5.0)}
    # The cleaned workbook round-trips the limits.
    again, _ = ingest_workbook(BytesIO(export_cleaned_workbook_bytes(result, project_metadata={})))
    assert chemistry_thresholds_from_readings(again.environmental_readings) == {"Benzene": (0.5, 5.0)}


def test_yellow_max_spelling_and_blank_cells_are_accepted() -> None:
    from ingestion import ingest_workbook

    collars = pd.DataFrame([{"hole_id": "BH-01", "easting": 0, "northing": 0, "elevation": 100, "total_depth": 10}])
    lith = pd.DataFrame([{"hole_id": "BH-01", "from_depth": 0, "to_depth": 10, "lithology_code": "Clay"}])
    env = pd.DataFrame(
        [
            {"hole_id": "BH-01", "parameter": "Chloride", "value": 12.0, "depth": 2.0, "green_max": 100, "yellow_max": 250},
            {"hole_id": "BH-01", "parameter": "Chloride", "value": 1.0, "depth": 4.0, "green_max": None, "yellow_max": None},
        ]
    )
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        collars.to_excel(writer, sheet_name="Collars", index=False)
        lith.to_excel(writer, sheet_name="Lithology", index=False)
        env.to_excel(writer, sheet_name="Environmental", index=False)
    buffer.seek(0)
    result, _ = ingest_workbook(buffer)
    assert not result.errors
    assert [(r.green_max, r.orange_max) for r in result.environmental_readings] == [(100.0, 250.0), (None, None)]


# --- export metadata ------------------------------------------------------------------


def test_metadata_payload_and_docx_list_the_limits() -> None:
    from docx_export import _metadata_text
    from ui_helpers import export_metadata_payload

    payload = export_metadata_payload(
        section_title="A-A'",
        preset_label=None,
        vertical_exaggeration=5.0,
        hole_count=3,
        transect_label="A-A'",
        overlap_warnings=(),
        chemistry_color_mode="threshold",
        chemistry_thresholds=THRESHOLDS,
    )
    assert payload["chemistry_label_colour_mode"] == "threshold"
    assert payload["chemistry_thresholds"] == {
        "Chloride": {"green_max": 100.0, "orange_max": 250.0},
        "Benzene": {"green_max": 0.5, "orange_max": 5.0},
    }
    json.dumps(payload)
    assert _metadata_text(payload["chemistry_thresholds"]) == (
        "Chloride: green_max 100, orange_max 250\nBenzene: green_max 0.5, orange_max 5"
    )
    # Black mode lists no limits.
    black = export_metadata_payload(
        section_title="A-A'", preset_label=None, vertical_exaggeration=None, hole_count=1,
        transect_label=None, overlap_warnings=(), chemistry_color_mode="black", chemistry_thresholds=THRESHOLDS,
    )
    assert "chemistry_thresholds" not in black


# --- Configure editor seeding -------------------------------------------------------------


def test_seed_rules_chloride_only_for_chloride_and_workbook_limits_win() -> None:
    from app_configure import _is_chloride_like, _limit_step, seeded_parameter_thresholds

    assert all(_is_chloride_like(n) for n in ("Chloride", "chloride (Cl)", "Cl", "Cl-", " CHLORIDE "))
    # Chlorinated solvents are not chloride: they must not start on 100 / 250 mg/kg.
    assert not any(_is_chloride_like(n) for n in ("Vinyl chloride", "Trichloroethene", "Chlorobenzene", "Clay"))
    assert seeded_parameter_thresholds("Chloride", {}, default_green=7.0, default_orange=9.0) == (100.0, 250.0)
    assert seeded_parameter_thresholds("Vinyl chloride", {}, default_green=7.0, default_orange=9.0) == (7.0, 9.0)
    assert seeded_parameter_thresholds("TPH", {}, default_green=None, default_orange=None) == (100.0, 250.0)
    # Workbook columns win, limit by limit; a lone green limit never sits above orange.
    assert seeded_parameter_thresholds("benzene", {"Benzene": (0.5, 5.0)}, default_green=7.0, default_orange=9.0) == (0.5, 5.0)
    assert seeded_parameter_thresholds("Chloride", {"Chloride": (300.0, None)}, default_green=None, default_orange=None) == (300.0, 300.0)
    # Spinner step follows the limits' magnitude.
    assert _limit_step(100.0, 250.0) == 1.0
    assert _limit_step(0.5, 5.0) == 0.1
    assert _limit_step(0.005, 0.05) == 0.001
    assert _limit_step(None, None) == 0.001


# --- Configure editor (AppTest) ----------------------------------------------------------


def test_configure_seeds_per_parameter_limits_and_remembers_them() -> None:
    from streamlit.testing.v1 import AppTest

    from app_configure import chemistry_threshold_widget_key

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=180)
    at.run()
    at.file_uploader[0].upload("two_params.xlsx", _two_parameter_workbook_bytes()).run()
    assert not at.exception
    assert at.session_state["parse_result"] is not None
    assert sorted(at.session_state["environmental_parameters_multiselect"]) == ["Benzene", "Chloride"]

    at.radio(key="chemistry_color_mode_radio").set_value("Green / orange / red thresholds").run()
    assert not at.exception
    token = str(at.session_state["file_hash"])
    keys = {
        name: (
            chemistry_threshold_widget_key(name, "green", token),
            chemistry_threshold_widget_key(name, "orange", token),
        )
        for name in ("Chloride", "Benzene")
    }
    inputs = {widget.key: widget for widget in at.number_input}
    for green_key, orange_key in keys.values():
        assert green_key in inputs and orange_key in inputs
    # Chloride-like parameters start on the client bands; benzene on its workbook limits.
    assert (inputs[keys["Chloride"][0]].value, inputs[keys["Chloride"][1]].value) == (100.0, 250.0)
    assert (inputs[keys["Benzene"][0]].value, inputs[keys["Benzene"][1]].value) == (0.5, 5.0)
    # Help text names the parameter.
    assert "Benzene" in (inputs[keys["Benzene"][0]].help or "")

    # Editing one parameter is remembered across reruns and does not touch the other.
    at.number_input(key=keys["Chloride"][0]).set_value(80.0).run()
    at.number_input(key=keys["Benzene"][1]).set_value(4.0).run()
    at.run()
    assert not at.exception
    assert at.session_state[keys["Chloride"][0]] == 80.0
    assert at.session_state[keys["Benzene"][0]] == 0.5

    # Deselecting a parameter removes its row; the other keeps its edited value.
    at.multiselect(key="environmental_parameters_multiselect").set_value(["Chloride"]).run()
    assert not at.exception
    inputs = {widget.key: widget for widget in at.number_input}
    assert keys["Benzene"][0] not in inputs
    assert inputs[keys["Chloride"][0]].value == 80.0
    # Re-selecting it brings back its edited limits, not the workbook seed.
    at.multiselect(key="environmental_parameters_multiselect").set_value(["Benzene", "Chloride"]).run()
    assert not at.exception
    inputs = {widget.key: widget for widget in at.number_input}
    assert (inputs[keys["Benzene"][0]].value, inputs[keys["Benzene"][1]].value) == (0.5, 4.0)
    # Small limits keep their spelling (0.5, not 0.50) and a finer spinner step.
    assert inputs[keys["Benzene"][0]].format == "%g" and inputs[keys["Benzene"][0]].step == 0.1
    assert inputs[keys["Chloride"][0]].step == 1.0

    # Leaving threshold mode hides the editor; coming back restores the edits.
    at.radio(key="chemistry_color_mode_radio").set_value("All black").run()
    assert not at.exception
    assert not any(widget.key == keys["Chloride"][0] for widget in at.number_input)
    at.radio(key="chemistry_color_mode_radio").set_value("Green / orange / red thresholds").run()
    inputs = {widget.key: widget for widget in at.number_input}
    assert inputs[keys["Chloride"][0]].value == 80.0

    # A new workbook starts afresh: its own keys, seeded again from the defaults.
    at.file_uploader[0].upload("two_params_b.xlsx", _two_parameter_workbook_bytes("B-B'")).run()
    assert not at.exception
    new_token = str(at.session_state["file_hash"])
    assert new_token != token
    at.radio(key="chemistry_color_mode_radio").set_value("Green / orange / red thresholds").run()
    assert not at.exception
    inputs = {widget.key: widget for widget in at.number_input}
    new_green = chemistry_threshold_widget_key("Chloride", "green", new_token)
    assert new_green in inputs and inputs[new_green].value == 100.0
    assert keys["Chloride"][0] not in inputs
