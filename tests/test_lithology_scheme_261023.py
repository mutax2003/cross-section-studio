"""Client lithology legend 261023 is the single source of lithology colours and hatches.

The expected table below was read from the 261023 legend PDF's vector fills
and hatch strokes (261023_Cross_Section_Litho_Legend.pdf) and agrees with the
CAD source (Cross_Section_Litho_Legend_261002_CB.dwg layer TrueColors / HATCH
patterns) and the workbook ``data/BH Log Lithology Legend.xlsx``.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import build_lithology_legend  # noqa: E402

import constants  # noqa: E402
from constants import (  # noqa: E402
    HATCH_GRAVEL,
    HATCH_PLUS,
    HATCH_SANDY,
    HATCH_SILTY,
    LITHOLOGY_LEGEND_ORDER,
    LITHOLOGY_SCHEME_ID,
    USGS_LITHOLOGY_COLORS,
    get_lithology_style,
)
from paths import bh_log_lithology_legend_path, bh_log_lithology_legend_xlsx_path  # noqa: E402

SOLID = ""

# Swatch -> (fill, hatch) on the 261023 legend sheet, in sheet order.
LEGEND_261023: dict[str, tuple[str, str]] = {
    "Clay": ("#967259", SOLID),
    "Silt": ("#8D5524", SOLID),
    "Sand": ("#FFE39F", HATCH_SANDY),
    "Loamy Sand": ("#FFE39F", HATCH_PLUS),
    "Sandy Clay": ("#967259", HATCH_SANDY),
    "Silty Clay": ("#967259", HATCH_SILTY),
    "Organics": ("#38220F", SOLID),
    "Topsoil": ("#534230", SOLID),
    "Loam": ("#C68642", HATCH_PLUS),
    "Silty Loam": ("#C68642", HATCH_SILTY),
    "Clay Loam": ("#CAB5A5", SOLID),
    "Sandy Clay Loam": ("#CAB5A5", HATCH_SANDY),
    "Silty Clay Loam": ("#CAB5A5", HATCH_SILTY),
    "Fill": ("#854442", SOLID),
    "Gravel": ("#D9D9D9", HATCH_GRAVEL),
    "Sand and Gravel": ("#FFE39F", HATCH_GRAVEL),
    "Drilling Waste": ("#808080", SOLID),
    "Other": ("#4D5D53", SOLID),
    "No Recovery": ("#FFFFFF", SOLID),
    "Bentonite": ("#BFBFBF", SOLID),
    "Coal": ("#000000", SOLID),
    "Refuse": ("#8C973D", SOLID),
    "Siltstone": ("#4C516D", HATCH_SILTY),
    "Sandstone": ("#4C516D", HATCH_SANDY),
    "Mudstone": ("#4C516D", SOLID),
}


def test_hatch_family_characters() -> None:
    assert (HATCH_SANDY, HATCH_SILTY, HATCH_PLUS, HATCH_GRAVEL) == (".", "/", "+", "O")


def test_scheme_is_261023_and_lists_every_legend_lithology_in_order() -> None:
    assert LITHOLOGY_SCHEME_ID == "261023"
    assert list(LITHOLOGY_LEGEND_ORDER) == list(LEGEND_261023)
    assert len(LEGEND_261023) == 25  # 24 workbook rows + Topsoil


@pytest.mark.parametrize("code", list(LEGEND_261023))
def test_every_legend_lithology_draws_its_261023_colour_and_hatch(code: str) -> None:
    colour, hatch = LEGEND_261023[code]
    style = get_lithology_style(code)
    assert (style.color.upper(), style.hatch) == (colour, hatch)
    # Workbook spellings in any case resolve to the same style.
    assert get_lithology_style(code.upper()) == style


def test_committed_json_matches_the_build_script(tmp_path: Path) -> None:
    """Drift guard: the JSON is generated from the workbook, never hand-edited."""
    expected = build_lithology_legend.render_json(build_lithology_legend.build_scheme())
    committed = bh_log_lithology_legend_path().read_text(encoding="utf-8")
    assert committed == expected, "run: python scripts/build_lithology_legend.py"
    assert build_lithology_legend.main(["--check"]) == 0
    stale = tmp_path / "stale.json"
    stale.write_text("{}\n", encoding="utf-8")
    assert build_lithology_legend.main(["--check", "--output", str(stale)]) == 1


def test_workbook_group_colour_fills_down_and_hatch_text_maps_to_families() -> None:
    rows = {row["name"]: row for row in build_lithology_legend.read_workbook_rows(bh_log_lithology_legend_xlsx_path())}
    assert len(rows) == 24 and "Topsoil" not in rows
    # Rows without a hex inherit the group colour above them.
    assert rows["Sandy Clay"]["color"] == rows["Silty Clay"]["color"] == "#967259"
    assert rows["Sandy Clay Loam"]["color"] == rows["Clay Loam"]["color"] == "#CAB5A5"
    assert rows["Sandstone"]["color"] == rows["Mudstone"]["color"] == "#4C516D"
    assert {row["hatch_family"] for row in rows.values()} == {"solid", "sandy", "silty", "loam", "gravel"}
    for name, row in rows.items():
        family_hatch = build_lithology_legend.HATCH_FAMILIES[row["hatch_family"]]["hatch"]
        assert (row["color"], family_hatch) == LEGEND_261023[name], name


def test_hatch_text_parsing_and_unknown_hatch_is_an_error() -> None:
    family = build_lithology_legend.hatch_family
    assert [family(t) for t in ("solid", "Sand hatch .", "Silty hatch /", "Loam hatch +", "Gravel hatch")] == [
        "solid",
        "sandy",
        "silty",
        "loam",
        "gravel",
    ]
    assert family(None) == "solid"
    with pytest.raises(ValueError, match="Unknown hatch"):
        family("Crosshatch x")


def test_workbook_with_a_conflicting_cell_fill_is_rejected(tmp_path: Path) -> None:
    import openpyxl
    from openpyxl.styles import PatternFill

    book = openpyxl.load_workbook(bh_log_lithology_legend_xlsx_path())
    sheet = book.worksheets[0]
    sheet["G5"].fill = PatternFill("solid", fgColor="FF123456")  # Sandy Clay row
    broken = tmp_path / "legend.xlsx"
    book.save(broken)
    with pytest.raises(ValueError, match="disagrees with group colour"):
        build_lithology_legend.build_scheme(broken)


def test_legacy_codes_keep_their_styles_but_stay_off_the_legend() -> None:
    payload = json.loads(bh_log_lithology_legend_path().read_text(encoding="utf-8"))
    legacy = payload["legacy_codes"]
    assert set(legacy) == {"Silty Sand", "Bedrock", "Limestone", "Shale", "Flare Pit Material", "Sand and Clay"}
    assert not set(legacy) & set(LITHOLOGY_LEGEND_ORDER)
    for code in legacy:
        assert code in USGS_LITHOLOGY_COLORS
    assert get_lithology_style("Silty Sand").color == "#FFE39F"
    assert get_lithology_style("Silty Sand").hatch == HATCH_SILTY
    assert get_lithology_style("Bedrock").hatch == "xxx"


def test_unknown_codes_keep_the_grey_fallback() -> None:
    style = get_lithology_style("Glacial Till (unlisted)")
    assert (style.color, style.hatch) in constants.FALLBACK_LITHOLOGY_STYLES


def test_no_lithology_hex_is_hard_coded_outside_the_scheme() -> None:
    """Colours live in the generated JSON only (constants, legend sheet, swatch export)."""
    hexes = {colour for colour, _hatch in LEGEND_261023.values()} - {"#000000", "#FFFFFF"}
    for relative in (
        "constants.py",
        "hatch_patterns.py",
        "lithology_legend_sheet.py",
        "scripts/export_lithology_swatches.py",
        "workbook_template.py",
    ):
        text = (ROOT / relative).read_text(encoding="utf-8").upper()
        assert not [h for h in hexes if h in text], relative


def test_old_workbook_names_still_map_through_aliases() -> None:
    from ai_quality import load_lithology_aliases, normalize_lithology_code

    aliases = load_lithology_aliases(str(ROOT / "data" / "lithology_aliases.json"))
    for code in USGS_LITHOLOGY_COLORS:
        assert normalize_lithology_code(code.upper(), aliases) == code
    assert normalize_lithology_code("silty clay loam", aliases) == "Silty Clay Loam"
    assert normalize_lithology_code("top soil", aliases) == "Topsoil"
    assert all(target in USGS_LITHOLOGY_COLORS for target in aliases.values())
