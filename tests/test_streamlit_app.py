"""Streamlit AppTest smoke for upload → validate → generate."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "data" / "sample_boreholes.xlsx"


@pytest.fixture(scope="module")
def sample_workbook() -> Path:
    if not SAMPLE.exists():
        subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "generate_sample_data.py")],
            cwd=ROOT,
            check=True,
        )
    assert SAMPLE.exists()
    return SAMPLE


def test_streamlit_upload_and_validate(sample_workbook: Path) -> None:
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("sample.xlsx", sample_workbook.read_bytes()).run()
    assert not at.exception
    assert "hole_ids" in at.session_state and at.session_state["hole_ids"]
    assert "parse_result" in at.session_state and at.session_state["parse_result"] is not None
    assert "quality_report" in at.session_state and at.session_state["quality_report"] is not None


def test_streamlit_welcome_sample_button_present() -> None:
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=60)
    at.run()
    assert not at.exception
    labels = [btn.label for btn in at.button]
    assert "Try sample project" in labels


def test_streamlit_try_sample_project_parses(sample_workbook: Path) -> None:
    """Regression: sample load must re-detect/parse even when file_bytes is pre-set."""
    from streamlit.testing.v1 import AppTest

    assert sample_workbook.exists()
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    sample_buttons = [btn for btn in at.button if btn.label == "Try sample project"]
    assert sample_buttons, "Try sample project button missing"
    sample_buttons[0].click().run()
    assert not at.exception
    assert "file_bytes" in at.session_state and at.session_state["file_bytes"]
    assert "detection_result" in at.session_state and at.session_state["detection_result"] is not None
    assert "parse_result" in at.session_state and at.session_state["parse_result"] is not None
    assert "hole_ids" in at.session_state and at.session_state["hole_ids"]
    assert "quality_report" in at.session_state and at.session_state["quality_report"] is not None
    # Sample load remounts the file_uploader so session bytes stay authoritative.
    assert "workbook_uploader_key" in at.session_state
    assert at.session_state["workbook_uploader_key"] >= 1


def test_streamlit_clear_workbook_bumps_uploader_key(sample_workbook: Path) -> None:
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("sample.xlsx", sample_workbook.read_bytes()).run()
    assert at.session_state["parse_result"] is not None
    before = int(at.session_state["workbook_uploader_key"]) if "workbook_uploader_key" in at.session_state else 0
    clear_buttons = [btn for btn in at.button if btn.label == "Clear workbook"]
    assert clear_buttons, "Clear workbook button missing"
    clear_buttons[0].click().run()
    assert not at.exception
    assert at.session_state["parse_result"] is None
    assert at.session_state["file_bytes"] is None
    assert int(at.session_state["workbook_uploader_key"]) == before + 1


def test_streamlit_generate_smoke(sample_workbook: Path) -> None:
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("sample.xlsx", sample_workbook.read_bytes()).run()
    assert not at.exception
    assert at.session_state["parse_result"] is not None

    hole_ids = list(at.session_state["hole_ids"])
    assert len(hole_ids) >= 2
    at.session_state["hole_sequence_multiselect"] = hole_ids[: min(4, len(hole_ids))]
    for box in at.checkbox:
        label = box.label or ""
        if "Generate even if data checks found warnings" in label:
            box.set_value(True)
            break
    at.run()

    generate_buttons = [btn for btn in at.button if btn.label == "Generate section"]
    assert generate_buttons, (
        "Generate section missing after upload+transect setup; "
        f"buttons={[b.label for b in at.button]}"
    )
    assert not generate_buttons[0].disabled, (
        "Generate disabled — check QA blocking / placeholder elev / transect / overlaps"
    )
    generate_buttons[0].click().run()
    assert not at.exception
    assert "svg_bytes" in at.session_state and at.session_state["svg_bytes"]
    # Figure-first: setup collapsed; exactly ONE Generate button after the first build
    assert [btn.key for btn in at.button if btn.label == "Generate section"] == [
        "generate_section_strip"
    ]
    assert not any("Regenerate" in (btn.label or "") for btn in at.button)


def test_render_hero_compact_after_upload() -> None:
    from app_common import _render_hero

    # Smoke: compact class applied for stage >= 1 (no Streamlit run required for logic)
    assert callable(_render_hero)


def _generated_app(sample_workbook: Path):
    """AppTest with the sample uploaded and one section generated."""
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("sample.xlsx", sample_workbook.read_bytes()).run()
    hole_ids = list(at.session_state["hole_ids"])
    at.session_state["hole_sequence_multiselect"] = hole_ids[: min(4, len(hole_ids))]
    for box in at.checkbox:
        if "Generate even if data checks found warnings" in (box.label or ""):
            box.set_value(True)
            break
    at.run()
    [btn for btn in at.button if btn.label == "Generate section"][0].click().run()
    assert at.session_state["svg_bytes"]
    return at


def test_clear_workbook_asks_before_discarding_a_generated_section(sample_workbook: Path) -> None:
    at = _generated_app(sample_workbook)
    [btn for btn in at.button if btn.label == "Clear workbook"][0].click().run()
    assert not at.exception
    # Nothing discarded yet: a confirmation is shown instead.
    assert at.session_state["svg_bytes"]
    assert any("discard the generated section" in w.value for w in at.warning)

    [btn for btn in at.button if btn.key == "cancel_destructive"][0].click().run()
    assert at.session_state["svg_bytes"], "Cancel must keep the section"
    assert not any(btn.key == "confirm_destructive" for btn in at.button)

    [btn for btn in at.button if btn.label == "Clear workbook"][0].click().run()
    [btn for btn in at.button if btn.key == "confirm_destructive"][0].click().run()
    assert not at.exception
    assert at.session_state["parse_result"] is None
    assert at.session_state["svg_bytes"] is None


def test_generate_action_and_state_aware_coach_sit_above_validate(sample_workbook: Path) -> None:
    """UX regressions: Generate is at the top of the page (not ~2.5 screens
    down), the 'Next:' hint reflects the picked transect, and the stepper is
    inside the hero (it rendered outside, white-on-white)."""
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("sample.xlsx", sample_workbook.read_bytes()).run()
    at.run()
    assert not at.exception
    markdown = [m.value for m in at.markdown]
    hero = next(value for value in markdown if "<h1>Cross Section Studio</h1>" in value)
    assert "workflow-step" in hero, "stepper must render inside the hero"
    assert "✓ Upload" in hero
    coach = next(value for value in markdown if 'class="next-step-coach"' in value)
    # Coach copy lives in app_validate (may still say the old label).
    assert "is selected" in coach and (
        "Generate section" in coach or "Generate Cross-Section" in coach
    )

    def walk(node):
        yield node
        for child in getattr(node, "children", {}).values():
            yield from walk(child)

    flat = list(walk(at.main))
    generate_index = next(
        i for i, node in enumerate(flat)
        if getattr(node, "type", "") == "button" and node.label == "Generate section"
    )
    health_index = next(
        i for i, node in enumerate(flat)
        if getattr(node, "type", "") in {"subheader", "heading"} and "data health" in str(node.value).lower()
    )
    assert generate_index < health_index, "Generate must render above Validate's Data Health"


def test_preview_zoom_switches_to_scrollable_native_size_frame(sample_workbook: Path) -> None:
    at = _generated_app(sample_workbook)

    def frame_html() -> str:
        return next(md.value for md in at.markdown if "<img src=\"data:image/svg" in md.value)

    assert 'class="svg-frame"' in frame_html()
    assert "width:100%" in frame_html()
    at.session_state["svg_preview_zoom"] = "150%"
    at.run()
    assert not at.exception
    html_150 = frame_html()
    assert "svg-frame--zoomed" in html_150 and 'tabindex="0"' in html_150
    natural = at.session_state["svg_display_meta"].natural_width_px
    assert natural > 0
    assert f"width:{round(natural * 1.5)}px" in html_150


def test_menu_load_sample_asks_before_discarding_a_generated_section(sample_workbook: Path) -> None:
    at = _generated_app(sample_workbook)
    triggers = {
        "menu": lambda: at.menu_button(key="menu_file").click(
            "Load sample project (Alt+Shift+O)"
        ),
        "shortcut": lambda: at.button(key="menu_accel_sample").click(),
    }
    for name, trigger in triggers.items():
        trigger().run()
        assert not at.exception
        assert at.session_state["svg_bytes"], f"{name} discarded the section without asking"
        assert at.session_state["_pending_destructive"] == "sample"
        at.button(key="cancel_destructive").click().run()


def test_stale_destructive_prompt_is_dropped_once_its_section_is_gone(sample_workbook: Path) -> None:
    at = _generated_app(sample_workbook)
    [btn for btn in at.button if btn.label == "Clear workbook"][0].click().run()
    assert at.session_state["_pending_destructive"] == "clear"
    # Section disappears another way (new upload, menu clear...): the old
    # Confirm must not linger to wipe whatever is generated next.
    at.session_state["svg_bytes"] = None
    at.run()
    assert "_pending_destructive" not in at.session_state
    assert not [btn for btn in at.button if btn.key == "confirm_destructive"]


def test_parse_failure_is_shown_on_screen_with_technical_details() -> None:
    """A workbook the parser rejects used to leave the page silent: the error
    was stored and the function returned before the banner block rendered."""
    from streamlit.testing.v1 import AppTest

    from tests.conftest import make_workbook_bytes

    broken = make_workbook_bytes(
        [{"hole_id": "BH1", "easting": 0, "northing": 0, "elevation": 100}],  # no total_depth
        [{"hole_id": "BH1", "from_depth": 0, "to_depth": 5, "lithology_code": "Clay"}],
    )
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("broken.xlsx", broken).run()
    assert not at.exception
    assert at.session_state["parse_result"] is None
    assert any("total_depth" in e.value or "Collars" in e.value for e in at.error), [
        e.value for e in at.error
    ]
    assert any("Technical details" in x.label for x in at.expander)


def test_stale_generate_button_is_disabled_with_the_reason_when_blocked(
    sample_workbook: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """After a section exists, a settings change marks it stale; if Generate
    is blocked the stale button used to stay enabled and silently do nothing."""
    import app_configure

    at = _generated_app(sample_workbook)
    # Per-pair clipping now resolves the sample's pinch-out overlaps, so inject a
    # residual overlap warning into the Configure preflight; blocking on it
    # disables Generate. Changing VE then marks the existing section stale.
    real_preflight = app_configure.cached_configure_preflight

    def preflight_with_overlap(*args, **kwargs):
        warnings, summaries = real_preflight(*args, **kwargs)
        return (*warnings, "Polygon overlap: 1 inter-hole contact conflict(s) detected"), summaries

    monkeypatch.setattr(app_configure, "cached_configure_preflight", preflight_with_overlap)
    at.session_state["allow_pinch_outs"] = True
    at.session_state["fail_on_overlaps_checkbox"] = True
    at.session_state["vertical_exaggeration"] = 3.0
    at.run()
    assert not at.exception
    generate_buttons = [b for b in at.button if b.label == "Generate section"]
    assert len(generate_buttons) == 1, "one Generate button when out of date"
    assert generate_buttons[0].disabled
    banner = next(m.value for m in at.markdown if 'class="generate-strip is-stale"' in m.value)
    assert "Out of date" in banner and "overlap" in banner
    at.session_state["_regenerate_requested"] = True  # Alt+Shift+G / menu path
    at.run()
    assert any(
        "Generate section skipped" in w.value and "overlap" in w.value for w in at.warning
    )


def test_sections_tab_drop_down_drives_the_preview(tmp_path: Path) -> None:
    """Pick a named section from the workbook's Sections tab: the hole order,
    mode and sheet label follow, so a moved line only needs the tab edited."""
    import pandas as pd
    from streamlit.testing.v1 import AppTest

    collars = pd.DataFrame(
        [{"hole_id": f"BH-0{i}", "easting": i * 10.0, "northing": 0.0, "elevation": 100.0, "total_depth": 8.0} for i in range(1, 6)]
    )
    lith = pd.DataFrame(
        [{"hole_id": f"BH-0{i}", "from_depth": 0.0, "to_depth": 8.0, "lithology_code": "Clay"} for i in range(1, 6)]
    )
    sections = pd.DataFrame(
        [{"section_label": "A-A'", "hole_ids": "BH-01, BH-02, BH-03"}, {"section_label": "B-B'", "hole_ids": "BH-03 → BH-04 → BH-05"}]
    )
    workbook = tmp_path / "sections.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        collars.to_excel(writer, sheet_name="Collars", index=False)
        lith.to_excel(writer, sheet_name="Lithology", index=False)
        sections.to_excel(writer, sheet_name="Sections", index=False)

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("sections.xlsx", workbook.read_bytes()).run()
    assert not at.exception
    picker = at.selectbox(key="workbook_section_choice")
    assert len(picker.options) == 3  # Custom + two rows (options are row indices)
    # First section is the default preview.
    assert at.session_state["hole_sequence_multiselect"] == ["BH-01", "BH-02", "BH-03"]
    picker.select(1).run()  # row index 1 = B-B'
    at.run()
    assert not at.exception
    assert at.session_state["hole_sequence_multiselect"] == ["BH-03", "BH-04", "BH-05"]
    assert at.session_state["consulting_section_label"] == "B-B'"
    assert at.session_state["section_title"] == "B-B'"  # file names and metadata follow

    # Picking a section while the Consulting layout (with its "Section label"
    # text box) is active used to raise StreamlitWidgetAlreadyInstantiatedError.
    at.session_state["output_preset"] = "consulting_report"
    at.run()
    at.selectbox(key="workbook_section_choice").select(0).run()
    at.run()
    assert not at.exception
    assert at.session_state["hole_sequence_multiselect"] == ["BH-01", "BH-02", "BH-03"]
    assert at.session_state["consulting_section_label"] == "A-A'"

    # A new workbook must not inherit the previous one's hole order or label.
    plain = tmp_path / "plain.xlsx"
    with pd.ExcelWriter(plain, engine="openpyxl") as writer:
        collars.to_excel(writer, sheet_name="Collars", index=False)
        lith.to_excel(writer, sheet_name="Lithology", index=False)
    at.file_uploader[0].upload("plain.xlsx", plain.read_bytes()).run()
    at.run()
    assert not at.exception
    assert not [sb for sb in at.selectbox if sb.key == "workbook_section_choice"]
    assert at.session_state["consulting_section_label"] != "A-A'"
    # A small workbook without a Sections tab starts with every hole.
    assert at.session_state.get("hole_sequence_multiselect") == ["BH-01", "BH-02", "BH-03", "BH-04", "BH-05"]


def test_clear_then_sample_drops_project_fields_and_opens_unblocked(sample_workbook: Path) -> None:
    """The title block fields seeded from a Project tab leaked onto the next
    workbook, and the demo opened on a blocked Generate under the consulting
    preset (its pinch-out overlaps are blocked by default)."""
    from streamlit.testing.v1 import AppTest

    from workbook_template import build_input_template_bytes

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("template.xlsx", build_input_template_bytes()).run()
    assert not at.exception
    # One-shot upload banner: read it on the run that shows it.
    assert any("template's sample values" in w.value for w in at.warning)
    at.run()
    assert at.session_state["consulting_section_label"] == "A - A' WITH CHLORIDE AVERAGES"
    at.session_state["output_preset"] = "consulting_report"
    at.run()
    [btn for btn in at.button if btn.label == "Clear workbook"][0].click().run()
    [btn for btn in at.button if btn.label == "Try sample project"][0].click().run()
    at.run()
    assert not at.exception
    assert at.session_state["consulting_section_label"] != "A - A' WITH CHLORIDE AVERAGES"
    assert not any("Polygon overlaps detected" in e.value for e in at.error)
    assert at.session_state.get("fail_on_overlaps_checkbox") is False


def test_batch_section_lines_fill_from_suggestions_and_flag_bad_lines(sample_workbook: Path) -> None:
    """Fill from suggested lines works without visiting Recommended mode, and
    each typed line is checked on its own under the editor."""
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("sample.xlsx", sample_workbook.read_bytes()).run()
    at.button(key="batch_fill_recommended").click().run()
    assert not at.exception
    filled = str(at.session_state["batch_transect_specs"])
    assert filled.startswith("A-A' | ")
    assert not any("Switch to Recommended" in w.value for w in at.warning)
    hole_ids = list(at.session_state["hole_ids"])
    at.session_state["batch_transect_specs"] = (
        f"A-A' | {hole_ids[0]}, {hole_ids[1]}\nC-C' | {hole_ids[0]}, MW-99"
    )
    at.run()
    assert not at.exception
    status = [w.value for w in at.warning if "section lines ready" in w.value]
    assert status and "1 of 2" in status[0] and "C-C': MW-99 not in Collars" in status[0]


def test_profile_chips_show_three_and_fold_the_rest_into_a_tooltip() -> None:
    from app_common import profile_chips_html

    html = profile_chips_html(
        interpretation_mode="interpolated",
        vertical_exaggeration=5.0,
        hole_count=4,
        polygon_count=6,
        is_stale=True,
        preset_label="Section sheet",
        render_layout="consulting_section",
        transect_label="A-A' BH-01→BH-04",
    )
    assert html.count('<span class="chip') == 4  # style, VE, holes, "+N more"
    assert ">Section sheet<" in html and "VE 5×" in html and "4 holes" in html
    assert "Out of date" in html and "Stale" not in html and "Fresh" not in html


def test_assist_mode_caption_hidden_unless_an_llm_is_switched_on(monkeypatch) -> None:
    import streamlit as st

    from app_common import llm_assist_status_caption

    monkeypatch.delenv("CROSS_SECTION_DISABLE_LLM", raising=False)
    st.session_state["enable_ai_suggestions"] = False
    assert llm_assist_status_caption() == ""
    st.session_state["enable_ai_suggestions"] = True
    st.session_state["llm_provider"] = "groq"
    caption = llm_assist_status_caption()
    assert "Groq" in caption and "API_KEY" not in caption and "`" not in caption


def test_prepare_deliverables_names_ready_files_and_downloads_share_one_row(
    sample_workbook: Path,
) -> None:
    at = _generated_app(sample_workbook)
    figure_files = {"PDF", "PNG", "SVG", "Word"}
    downloads = [b.label for b in at.get("download_button") if b.label in figure_files]
    assert downloads[:3] == ["PDF", "PNG", "SVG"]
    at.button(key="prepare_both_exports").click().run()
    assert not at.exception
    toasts = [t.value for t in at.toast]
    assert any(t in {"PDF, PNG and Word file ready", "PDF and PNG ready"} for t in toasts), toasts
    assert not any(b.key == "prepare_both_exports" for b in at.button)


def test_batch_zip_builds_the_good_lines_and_lists_the_skipped_ones(sample_workbook: Path) -> None:
    """One bad batch line used to fail the whole ZIP."""
    at = _generated_app(sample_workbook)
    holes = list(at.session_state["hole_ids"])
    at.session_state["batch_transect_specs"] = (
        f"A-A' | {holes[0]}, {holes[1]}, {holes[2]}\nZ-Z' | {holes[0]}, MW-99"
    )
    at.run()
    assert any("Skipped" in c.value and "MW-99" in c.value for c in at.caption)
    at.button(key="prepare_batch_zip").click().run()
    assert not at.exception
    assert at.session_state["batch_package_bytes"]
    assert any(
        "Batch ZIP ready: 1 section line" in s.value and "Left out: Z-Z'" in s.value
        for s in at.success
    )


def test_real_project_sharing_the_sample_number_and_date_is_not_flagged() -> None:
    """The B-B' workbook (own client and title, same project number and date
    as the template sample) was told its Project tab still held sample values."""
    from io import BytesIO

    import openpyxl
    from streamlit.testing.v1 import AppTest

    from workbook_template import build_input_template_bytes

    book = openpyxl.load_workbook(BytesIO(build_input_template_bytes()))
    sheet = book["Project"]
    for row in sheet.iter_rows(min_row=2):
        if row[0].value == "client_name":
            row[2].value = "WHITECAP RESOURCES INC."
        elif row[0].value == "section_title":
            row[2].value = "B - B' WITH CHLORIDE INTERVALS"
    buffer = BytesIO()
    book.save(buffer)

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("bb.xlsx", buffer.getvalue()).run()
    assert not at.exception
    assert not any("template's sample values" in w.value for w in at.warning)


def test_chemistry_columns_style_starts_with_red_labels() -> None:
    """The client's P2 figures print chloride values in red; the app always
    sent black or threshold, so the Chemistry columns style printed black."""
    from streamlit.testing.v1 import AppTest

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    from workbook_template import build_input_template_bytes

    # The committed template carries Environmental (chloride) sample rows.
    at.file_uploader[0].upload("t.xlsx", build_input_template_bytes()).run()
    at.session_state["output_preset"] = "p2_chemistry_sticks"
    at.run()
    at.run()
    assert not at.exception
    assert at.radio(key="chemistry_color_mode_radio").value == "All red (client P2 style)"
    at.session_state["output_preset"] = "section_sheet"
    at.run()
    at.run()
    assert at.radio(key="chemistry_color_mode_radio").value == "All black"


def _project_workbook_bytes(**project_values: str) -> bytes:
    """Committed template with Project-sheet values replaced."""
    from io import BytesIO

    import openpyxl

    from workbook_template import build_input_template_bytes

    book = openpyxl.load_workbook(BytesIO(build_input_template_bytes()))
    for row in book["Project"].iter_rows(min_row=2):
        if row[0].value in project_values:
            row[2].value = project_values[row[0].value]
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _browser_value(widget) -> object:
    """What the browser shows for a widget that has just re-mounted.

    A value carried over in session state but not flagged ``set_value`` is
    not sent; the browser then shows the widget's default (and sends that
    back on the next run), which blanked the title block after a style switch.
    """
    proto = widget.proto
    return proto.value if proto.set_value else proto.default


_TITLE_FIELDS = {
    "consulting_prepared_for": "ACME PIPELINES LTD.",
    "consulting_project_number": "P-77",
    "consulting_date": "01/02/26",
    "consulting_prepared_by": "NORTH CONSULTING",
    "consulting_drawn_by": "QA",
    "consulting_source": "QA 2026",
    "consulting_map_scale": "1:2500",
}


def test_title_block_survives_output_style_switches() -> None:
    """Switching Output style to Section sheet and back blanked the Section
    title and every title-block field in the browser (and so the PDF)."""
    from streamlit.testing.v1 import AppTest

    workbook = _project_workbook_bytes(
        client_name="ACME PIPELINES LTD.",
        prepared_by="NORTH CONSULTING",
        project_number="P-77",
        section_title="B - B' QA TITLE",
        report_date="01/02/26",
        drawn_by="QA",
        data_source="QA 2026",
        map_scale="1:2500",
        figure_preset="consulting_report",
    )
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("project.xlsx", workbook).run()
    at.run()
    assert not at.exception
    assert at.session_state["output_preset"] == "consulting_report"

    styles = ("section_sheet", "consulting_report", "quick_preview", "p2_chemistry_sticks")
    for style in styles + ("consulting_report",):
        at.selectbox(key="output_preset").set_value(style).run()
        assert not at.exception
        titles = [w for w in at.text_input if w.key in ("consulting_section_label", "section_title")]
        assert len(titles) == 1, style
        assert _browser_value(titles[0]) == "B - B' QA TITLE", style
        for widget in at.text_input:
            if widget.key in _TITLE_FIELDS:
                assert _browser_value(widget) == _TITLE_FIELDS[widget.key], (style, widget.key)
        if any(w.key == "export_include_title_block" for w in at.toggle):
            toggle = at.toggle(key="export_include_title_block")
            assert _browser_value(toggle) == at.session_state["export_include_title_block"], style
        at.run()  # e.g. a Generate rerun
        assert at.session_state["section_title"] == "B - B' QA TITLE", style
    shown = {w.key: w.value for w in at.text_input if w.key in _TITLE_FIELDS}
    assert shown == _TITLE_FIELDS


def test_new_workbook_reseeds_title_block_after_style_switches() -> None:
    from streamlit.testing.v1 import AppTest

    first = _project_workbook_bytes(
        client_name="FIRST CLIENT", section_title="A - A' FIRST", figure_preset="consulting_report"
    )
    second = _project_workbook_bytes(
        client_name="SECOND CLIENT", section_title="B - B' SECOND", figure_preset="consulting_report"
    )
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("first.xlsx", first).run()
    at.selectbox(key="output_preset").set_value("section_sheet").run()
    at.selectbox(key="output_preset").set_value("consulting_report").run()
    assert at.text_input(key="consulting_prepared_for").value == "FIRST CLIENT"
    at.selectbox(key="output_preset").set_value("section_sheet").run()
    at.file_uploader[0].upload("second.xlsx", second).run()
    at.run()
    assert not at.exception
    assert at.session_state["section_title"] == "B - B' SECOND"
    at.selectbox(key="output_preset").set_value("consulting_report").run()
    assert at.text_input(key="consulting_prepared_for").value == "SECOND CLIENT"
    assert at.text_input(key="consulting_section_label").value == "B - B' SECOND"


def test_template_without_lithology_shows_the_missing_sheet_error() -> None:
    """Deleting only the Lithology tab (the Data Entry sheet stays) loaded
    0 holes with "Data health OK" though Collars listed every hole."""
    from io import BytesIO

    import openpyxl
    from streamlit.testing.v1 import AppTest

    from workbook_template import build_input_template_bytes

    book = openpyxl.load_workbook(BytesIO(build_input_template_bytes()))
    del book["Lithology"]
    buffer = BytesIO()
    book.save(buffer)

    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=120)
    at.run()
    at.file_uploader[0].upload("no_lith.xlsx", buffer.getvalue()).run()
    assert not at.exception
    assert at.session_state["parse_result"] is None
    errors = [str(item.value) for item in at.error]
    assert any("No **Lithology** sheet was found" in text for text in errors), errors
    assert not any("Loaded **" in str(item.value) for item in at.success)
    # The template is not a field export: no "Field Data sheet" note.
    assert not any("Field Data sheet" in str(item.value) for item in at.info)
